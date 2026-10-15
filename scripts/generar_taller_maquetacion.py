# -*- coding: utf-8 -*-
"""Genera los materiales del Taller de maquetacion (Tema 2) y de la carta.

Salidas:
  web/descargas/t1-pedidos.docx          T1: datos separados por tabulaciones, sin tabulaciones fijadas
  web/descargas/t5-invitados.xlsx        T5: origen de datos de la combinacion con reglas
  web/descargas/carta-contenido.docx     Carta: textos sin formato del restaurante inventado
  web/assets/img/taller/t1-autorizacion.png   T1: aspecto final del modelo de autorizacion
  web/assets/img/taller/t2-tabla.png          T2: tabla que hay que reproducir
  web/assets/img/taller/t3-portada.svg        T3: portada de referencia
  web/assets/img/taller/t4-grafico.png        T4: grafico de anillo de referencia
  web/assets/img/taller/carta-boceto.svg      Carta: boceto de distribucion (no un modelo)

Todas las imagenes son propias: sustituyen a las del material original, que
llevaban marca de agua de terceros o reproducian la carta de un negocio real.
Todos los datos son inventados.

Las de T1 y T2 se obtienen renderizando con LibreOffice un .docx de
referencia construido aqui mismo, para que se parezcan a lo que sale en Word.

Uso:  python scripts/generar_taller_maquetacion.py   (necesita soffice y pdftoppm)
"""
from pathlib import Path
import subprocess
import tempfile
import math

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image

RAIZ = Path(__file__).resolve().parents[1]
DESCARGAS = RAIZ / 'web' / 'descargas'
IMG = RAIZ / 'web' / 'assets' / 'img' / 'taller'
IMG.mkdir(parents=True, exist_ok=True)


def num(valor, decimales=1):
    """Formato espanol: separador de miles con punto y decimales con coma."""
    texto = f'{valor:,.{decimales}f}'
    return texto.replace(',', 'X').replace('.', ',').replace('X', '.')


def documento_base():
    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for lado in ('left_margin', 'right_margin'):
        setattr(sec, lado, Cm(2.5))
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    return doc


def renderizar(doc, destino, recorte_px=40):
    """Guarda el docx, lo pasa a PDF con LibreOffice y recorta la primera pagina."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        origen = tmp / 'ref.docx'
        doc.save(origen)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(origen)],
                       check=True, capture_output=True)
        subprocess.run(['pdftoppm', '-r', '110', '-png', '-f', '1', '-l', '1', str(tmp / 'ref.pdf'), str(tmp / 'p')],
                       check=True)
        png = next(tmp.glob('p*.png'))
        imagen = Image.open(png).convert('RGB')
        # Recorta el blanco sobrante dejando un margen.
        gris = imagen.convert('L').point(lambda v: 0 if v > 245 else 255)
        caja = gris.getbbox()
        caja = (max(caja[0] - recorte_px, 0), max(caja[1] - recorte_px, 0),
                min(caja[2] + recorte_px, imagen.width), min(caja[3] + recorte_px, imagen.height))
        imagen.crop(caja).save(destino, optimize=True)


def fijar_anchos(tabla, anchos):
    """Anchos de columna que respetan tanto Word como LibreOffice (rejilla + celdas)."""
    tabla.autofit = False
    grid = tabla._tbl.tblGrid
    for col, ancho in zip(grid.findall(qn('w:gridCol')), anchos):
        col.set(qn('w:w'), str(int(ancho.twips)))
    for i, ancho in enumerate(anchos):
        for celda in tabla.columns[i].cells:
            celda.width = ancho


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


# ---------------------------------------------------------------- T1 pedidos
def t1_pedidos():
    doc = documento_base()
    p = doc.add_paragraph()
    r = p.add_run('PEDIDOS DEL MES · Jerte Digital')
    r.bold = True
    r.font.size = Pt(14)
    filas = [
        ('Cód. producto', 'Fecha de pedido', 'Precio', 'Unidades'),
        ('MN24FHD', '02/10/2026', '129,90 €', '4'),
        ('MN27QHD', '05/10/2026', '249,00 €', '2'),
        ('PT14I5', '07/10/2026', '1.049,95 €', '1'),
        ('TC-MEC01', '09/10/2026', '59,50 €', '12'),
        ('RT-PRO-X', '14/10/2026', '89,99 €', '6'),
        ('SSD1TB', '16/10/2026', '74,25 €', '15'),
        ('CBL-HDMI2', '21/10/2026', '8,40 €', '40'),
        ('IMP-LASER', '28/10/2026', '189,00 €', '3'),
    ]
    for i, fila in enumerate(filas):
        p = doc.add_paragraph('\t' + '\t'.join(fila))
        if i == 0:
            for run in p.runs:
                run.bold = True
    doc.add_paragraph()
    doc.add_paragraph('Las columnas están separadas con una tabulación cada una, pero todavía '
                      'no hay ninguna tabulación fijada en la regla: por eso no se alinean.')
    doc.save(DESCARGAS / 't1-pedidos.docx')


# ------------------------------------------------- T1 autorizacion (referencia)
def t1_autorizacion():
    doc = documento_base()

    def parrafo(texto='', paradas=(), alinear=None, negrita=False, tam=None, despues=10):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(despues)
        for pos, tipo, relleno in paradas:
            p.paragraph_format.tab_stops.add_tab_stop(Cm(pos), tipo, relleno)
        if alinear:
            p.alignment = alinear
        if texto:
            r = p.add_run(texto)
            r.bold = negrita
            if tam:
                r.font.size = Pt(tam)
        return p

    LINEA = WD_TAB_LEADER.LINES
    IZQ = WD_TAB_ALIGNMENT.LEFT
    parrafo('AUTORIZACIÓN PARA ACTIVIDADES COMPLEMENTARIAS', alinear=WD_ALIGN_PARAGRAPH.CENTER,
            negrita=True, tam=14, despues=2)
    parrafo('(a cumplimentar por la madre, el padre o la persona que ejerza la tutela legal '
            'del alumnado menor de edad)', alinear=WD_ALIGN_PARAGRAPH.CENTER, tam=10, despues=18)
    parrafo('D./D.ª\t, con DNI/NIE\t,', paradas=[(10.5, IZQ, LINEA), (15.8, IZQ, LINEA)])
    parrafo('en calidad de madre/padre/tutor/tutora legal de\t,', paradas=[(15.8, IZQ, LINEA)])
    parrafo('AUTORIZO a participar en la siguiente actividad:', despues=6)
    parrafo('\t', paradas=[(16, IZQ, LINEA)])
    parrafo('\t', paradas=[(16, IZQ, LINEA)], despues=18)
    parrafo('En\t, a\tde\tde 20\t.', paradas=[(6, IZQ, LINEA), (7.6, IZQ, LINEA), (11.6, IZQ, LINEA),
                                               (15.6, IZQ, LINEA)], despues=18)
    parrafo('Firma:', despues=4)
    tabla = doc.add_table(rows=1, cols=1)
    tabla.style = 'Table Grid'
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.rows[0].height = Cm(2.6)
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    fijar_anchos(tabla, [Cm(6)])
    parrafo('', despues=6)
    parrafo('Adjunto copia de mi DNI/NIE en vigor.', tam=10)
    renderizar(doc, IMG / 't1-autorizacion.png')


# ------------------------------------------------------------- T2 la tabla
T2_FILAS = [
    # indicador, real 2024, variacion 2024 frente a 2023 (%), objetivo 2025, real 2025, unidad
    ('Cereza cosechada', 24500, 6.2, 25700, 23900, 't'),
    ('Cereza vendida en fresco', 19800, 5.1, 21000, 19200, 't'),
    ('Ventas totales', 61.4, 9.8, 66.0, 63.1, 'm €'),
    ('Ventas en el extranjero', 18.7, 14.3, 21.0, 20.2, 'm €'),
    ('Venta directa y comercio en línea', 3.1, 22.5, 3.8, 3.6, 'm €'),
]


def t2_valor(valor, unidad, variacion):
    dec = 0 if unidad == 't' else 1
    signo = '+' if variacion >= 0 else '−'
    return f'{num(valor, dec)} {unidad}', f'({signo}{num(abs(variacion), 1)} %)'


def t2_tabla():
    doc = documento_base()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    VERDE_OSCURO, VERDE_CLARO, VERDE_OBJ = '375623', 'E2EFDA', 'A9D08E'
    tabla = doc.add_table(rows=3 + len(T2_FILAS), cols=5)
    tabla.style = 'Table Grid'
    anchos = [Cm(1.0), Cm(8.2), Cm(2.6), Cm(2.6), Cm(2.6)]
    fijar_anchos(tabla, anchos)

    def escribir(celda, lineas, negrita=False, color=None, alinear=WD_ALIGN_PARAGRAPH.LEFT, tam=10):
        celda.text = ''
        for i, linea in enumerate(lineas):
            p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
            p.alignment = alinear
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(linea)
            r.bold = negrita
            r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    # Fila 0: titulo combinado en las cinco columnas.
    titulo = tabla.cell(0, 0).merge(tabla.cell(0, 4))
    escribir(titulo, ['INDICADORES DE LA TEMPORADA DE LA CEREZA'], True, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER, 11)
    sombrear(titulo, VERDE_OSCURO)
    # Fila 1: cabeceras.
    cab = tabla.cell(1, 0).merge(tabla.cell(1, 1))
    escribir(cab, ['Indicador'], True)
    for col, texto in zip((2, 3, 4), ('Real 2024', 'Objetivo 2025', 'Real 2025')):
        escribir(tabla.cell(1, col), [texto], True, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    # Fila 2: grupo con fechas.
    grupo = tabla.cell(2, 0).merge(tabla.cell(2, 1))
    escribir(grupo, ['Producción y ventas'], True)
    for col in (2, 3, 4):
        escribir(tabla.cell(2, col), ['31/07/' + ('2024' if col == 2 else '2025')], True,
                 alinear=WD_ALIGN_PARAGRAPH.CENTER)
    for col in range(5):
        for fila in (1, 2):
            sombrear(tabla.cell(fila, col), VERDE_CLARO)
    # Datos.
    for i, (nombre, real24, var24, obj25, real25, unidad) in enumerate(T2_FILAS):
        fila = 3 + i
        escribir(tabla.cell(fila, 0), [str(i + 1)], alinear=WD_ALIGN_PARAGRAPH.CENTER)
        escribir(tabla.cell(fila, 1), [nombre])
        var_obj = (obj25 / real24 - 1) * 100
        var_r25 = (real25 / real24 - 1) * 100
        for col, (valor, var) in zip((2, 3, 4), ((real24, var24), (obj25, var_obj), (real25, var_r25))):
            escribir(tabla.cell(fila, col), list(t2_valor(valor, unidad, var)), alinear=WD_ALIGN_PARAGRAPH.RIGHT)
        sombrear(tabla.cell(fila, 3), VERDE_OBJ)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run('t: toneladas · m €: millones de euros · Datos inventados para la práctica.')
    r.italic = True
    r.font.size = Pt(9)
    renderizar(doc, IMG / 't2-tabla.png', recorte_px=25)


# ------------------------------------------------------------- T3 portada
def t3_portada():
    def pentagono(cx, cy, r, giro):
        pts = []
        for k in range(5):
            a = math.radians(giro + k * 72 - 90)
            pts.append(f'{cx + r * math.cos(a):.1f},{cy + r * math.sin(a):.1f}')
        return ' '.join(pts)

    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="-12 -12 234 321" role="img"
  aria-labelledby="t3t t3d">
<title id="t3t">Portada de referencia para la práctica T3</title>
<desc id="t3d">Página A4 vertical con dos triángulos granates y dos rosados arriba y abajo, dos pentágonos
verdes semitransparentes que se salen por las esquinas superior izquierda e inferior derecha, y en el centro
el texto PROYECTO curvado en arco con degradado, con el subtítulo Ciclo de SMR, curso 2026/2027.</desc>
<defs>
  <clipPath id="hoja"><rect x="0" y="0" width="210" height="297"/></clipPath>
  <linearGradient id="deg" x1="0" y1="0" x2="0" y2="1">
    <stop offset="0" stop-color="#C0392B"/><stop offset="1" stop-color="#5B0E14"/>
  </linearGradient>
  <path id="arco" d="M 8 172 Q 105 92 202 172"/>
</defs>
<rect x="-12" y="-12" width="234" height="321" fill="#B7B7B7"/>
<rect x="0" y="0" width="210" height="297" fill="#FFFFFF"/>
<g clip-path="url(#hoja)">
  <polygon points="0,0 210,0 0,105" fill="#F4C7CC"/>
  <polygon points="45,0 210,0 150,62" fill="#7B1420"/>
  <polygon points="210,297 0,297 210,192" fill="#F4C7CC"/>
  <polygon points="165,297 0,297 60,235" fill="#7B1420"/>
  <polygon points="{pentagono(22, 28, 44, 12)}" fill="#3A7D44" fill-opacity="0.6"/>
  <polygon points="{pentagono(190, 272, 44, 30)}" fill="#3A7D44" fill-opacity="0.6"/>
</g>
<rect x="0" y="0" width="210" height="297" fill="none" stroke="#222" stroke-width="0.8"/>
<text font-family="Arial Black, Arial, sans-serif" font-weight="900" font-size="31" fill="url(#deg)"
  text-anchor="middle"><textPath href="#arco" startOffset="50%">PROYECTO</textPath></text>
<text x="105" y="190" font-family="Arial, sans-serif" font-size="9" fill="#5B0E14" text-anchor="middle">Ciclo de SMR · Curso 2026/2027</text>
</svg>
'''
    (IMG / 't3-portada.svg').write_text(svg, encoding='utf-8')


# ------------------------------------------------------------- T4 grafico
T4_DATOS = [('Tienda física', 38.5), ('Web propia', 27.2), ('Marketplaces', 21.6), ('Teléfono', 12.7)]


def t4_grafico():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    assert abs(sum(v for _, v in T4_DATOS) - 100) < 0.01
    colores = ['#2E75B6', '#ED7D31', '#70AD47', '#FFC000']
    fig, ax = plt.subplots(figsize=(6.2, 3.6), dpi=130)
    cunas, _ = ax.pie([v for _, v in T4_DATOS], colors=colores, startangle=90, counterclock=False,
                      wedgeprops=dict(width=0.4, edgecolor='white', linewidth=1.5))
    for cuna, (_, valor) in zip(cunas, T4_DATOS):
        ang = math.radians((cuna.theta1 + cuna.theta2) / 2)
        x, y = math.cos(ang), math.sin(ang)
        ax.annotate(f'{num(valor, 2)} %', xy=(0.82 * x, 0.82 * y), xytext=(1.32 * x, 1.25 * y),
                    ha='center', va='center', fontsize=9,
                    arrowprops=dict(arrowstyle='-', color='#7F7F7F', linewidth=0.8))
    ax.legend(cunas, [n for n, _ in T4_DATOS], loc='center left', bbox_to_anchor=(1.12, 0.5),
              frameon=False, fontsize=9)
    ax.set_title('Ventas por canal · 2025', fontsize=12)
    ax.set_aspect('equal')
    fig.savefig(IMG / 't4-grafico.png', bbox_inches='tight', facecolor='white')
    plt.close(fig)


# ------------------------------------------------------------- T5 invitados
T5_INVITADOS = [
    # nombre, apellidos, correo, asiste, regalo
    ('Ana', 'Sánchez Gil', 'ana.sanchez@correo.example', 'Sí', 'Sí'),
    ('Bruno', 'Martín Calle', 'bruno.martin@correo.example', 'Sí', 'No'),
    ('Carmen', 'Vidal Ríos', 'carmen.vidal@correo.example', 'No', 'Sí'),
    ('Daniel', 'Pérez Moreno', 'daniel.perez@correo.example', 'Sí', 'Sí'),
    ('Elena', 'Robledo Paz', 'elena.robledo@correo.example', 'Pendiente', 'No'),
    ('Fernando', 'Iglesias Mora', 'fernando.iglesias@correo.example', 'Sí', 'No'),
    ('Gloria', 'Castro León', 'gloria.castro@correo.example', 'Sí', 'Sí'),
    ('Hugo', 'Navarro Sanz', 'hugo.navarro@correo.example', 'No', 'No'),
    ('Irene', 'Blanco Ruiz', 'irene.blanco@correo.example', 'Sí', 'No'),
    ('Javier', 'Ortega Díaz', 'javier.ortega@correo.example', 'Sí', 'Sí'),
    ('Laura', 'Romero Prieto', 'laura.romero@correo.example', 'Pendiente', 'Sí'),
    ('Mario', 'Serrano Vega', 'mario.serrano@correo.example', 'Sí', 'Sí'),
]


def t5_invitados():
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Invitados'
    hoja.append(['Nombre', 'Apellidos', 'Correo', 'Asiste', 'Regalo'])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in T5_INVITADOS:
        hoja.append(list(fila))
    for col, ancho in zip('ABCDE', (12, 18, 34, 11, 9)):
        hoja.column_dimensions[col].width = ancho
    libro.save(DESCARGAS / 't5-invitados.xlsx')


# ------------------------------------------------------------- Carta
CARTA = {
    'Para picar': [
        ('Jamón ibérico de bellota cortado a cuchillo', 'Con pan de pueblo tostado y tomate rallado', '19,50', '1'),
        ('Torta del Casar con picos', 'Para compartir entre dos o tres personas', '14,00', '1, 7'),
        ('Pimientos de Padrón', 'Con sal en escamas', '7,50', ''),
        ('Croquetas caseras de jamón', 'Seis unidades', '9,00', '1, 3, 7'),
        ('Patatas revolconas con torreznos', 'Con pimentón de la Vera', '8,50', ''),
    ],
    'Ensaladas': [
        ('Ensalada de cerezas y queso de cabra', 'Brotes tiernos, nueces y vinagreta de miel', '11,50', '7, 8, 12'),
        ('Ensalada de pimientos asados', 'Con ventresca de bonito y cebolleta', '10,50', '4'),
        ('Ensalada de la huerta', 'Tomate, lechuga, cebolla, aceitunas y espárragos', '8,00', ''),
    ],
    'Carnes': [
        ('Caldereta de cabrito', 'Receta de la abuela, guisada a fuego lento', '18,50', ''),
        ('Presa ibérica a la brasa', 'Con patatas panaderas y pimientos', '17,00', ''),
        ('Secreto ibérico', 'Con salsa de cereza picota', '16,50', '12'),
        ('Migas extremeñas', 'Con uvas, chorizo y huevo frito', '9,50', '1, 3'),
        ('Hamburguesa de ternera del Valle', 'Pan brioche, queso de cabra y cebolla caramelizada', '13,00', '1, 3, 7, 11'),
    ],
    'Pescados': [
        ('Trucha a la navarra', 'Rellena de jamón, con patatas', '14,50', '4'),
        ('Bacalao con pimentón de la Vera', 'Sobre pisto de verduras', '16,00', '4'),
        ('Pulpo a la brasa', 'Con cachelos y aceite de pimentón', '19,00', '14'),
    ],
    'Postres': [
        ('Tarta de cerezas del Jerte', 'Masa quebrada y cerezas en almíbar', '6,00', '1, 3, 7'),
        ('Picotas con nata', 'En temporada (de mayo a julio)', '5,50', '7'),
        ('Perrunillas caseras', 'Tres unidades, con café o infusión', '4,50', '1, 3'),
        ('Flan de queso de cabra', 'Con caramelo de miel', '5,50', '3, 7'),
    ],
    'Bebidas': [
        ('Agua mineral (50 cl)', '', '1,80', ''),
        ('Agua con gas (50 cl)', '', '2,20', ''),
        ('Refrescos', '', '2,50', ''),
        ('Mosto de cereza del Valle', 'Sin alcohol', '3,00', ''),
        ('Zumo natural de naranja', '', '3,50', ''),
        ('Café o infusión', '', '1,50', '7'),
    ],
}

ALERGENOS = ['Cereales con gluten', 'Crustáceos', 'Huevos', 'Pescado', 'Cacahuetes', 'Soja', 'Leche',
             'Frutos de cáscara', 'Apio', 'Mostaza', 'Granos de sésamo', 'Dióxido de azufre y sulfitos',
             'Altramuces', 'Moluscos']


def carta_contenido():
    doc = documento_base()
    doc.add_paragraph('LA CEREZA BRAVA')
    doc.add_paragraph('Cocina tradicional del Valle del Jerte')
    doc.add_paragraph('Restaurante inventado para la práctica. Todos los datos son ficticios.')
    doc.add_paragraph('')
    doc.add_paragraph('TEXTOS DEL EXTERIOR')
    for linea in ['Reservas: 927 123 456',
                  'Horario: de martes a domingo, de 13:00 a 16:00 y de 20:30 a 23:30. Lunes cerrado.',
                  'Dirección: Plaza de la Cereza, 3 · Valle del Jerte (Cáceres)',
                  'Redes: @lacerezabrava',
                  'Precios con IVA incluido. Pan y aperitivo: 1,50 € por persona.',
                  'Si tienes alguna alergia o intolerancia, avisa al personal antes de pedir.',
                  'ALÉRGENOS']:
        doc.add_paragraph(linea)
    for i, nombre in enumerate(ALERGENOS, 1):
        doc.add_paragraph(f'{i} | {nombre}')
    doc.add_paragraph('')
    doc.add_paragraph('TEXTOS DEL INTERIOR (formato de cada plato: Nombre | Precio | Alérgenos, y debajo la descripción)')
    for seccion, platos in CARTA.items():
        doc.add_paragraph(seccion.upper())
        for nombre, descripcion, precio, alergenos in platos:
            doc.add_paragraph(f'{nombre} | {precio} € | {alergenos}'.rstrip(' |'))
            if descripcion:
                doc.add_paragraph(descripcion)
    doc.save(DESCARGAS / 'carta-contenido.docx')


def carta_boceto():
    svg = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 620 470" role="img" aria-labelledby="cbt cbd">
<title id="cbt">Boceto de distribución de la carta</title>
<desc id="cbd">Dos hojas A4 apaisadas. Exterior: a la izquierda la contraportada con alérgenos, reservas,
horario y redes; a la derecha la portada con logotipo, nombre y una imagen. Interior: tres columnas con las
secciones de la carta y los precios alineados a la derecha.</desc>
<style>
 text{font-family:Segoe UI,Calibri,Arial,sans-serif;fill:#333}
 .t{font-size:13px;font-weight:700}.s{font-size:10.5px}.c{font-size:10px;fill:#666}
 .hoja{fill:#fff;stroke:#555;stroke-width:1.2}.zona{fill:#EEF3F9;stroke:#9DB3CC;stroke-dasharray:4 3}
 .pliegue{stroke:#999;stroke-dasharray:6 4}
</style>
<rect width="620" height="470" fill="#F7F7F7"/>
<text x="20" y="22" class="t">Hoja 1 · Exterior (A4 horizontal)</text>
<rect class="hoja" x="20" y="32" width="580" height="190"/>
<line class="pliegue" x1="310" y1="32" x2="310" y2="222"/>
<rect class="zona" x="34" y="46" width="262" height="78"/><text x="44" y="66" class="s">Alérgenos: 14 iconos numerados</text><text x="44" y="84" class="c">(formas o iconos + texto)</text>
<rect class="zona" x="34" y="134" width="262" height="74"/><text x="44" y="154" class="s">Reservas · horario · dirección · redes</text><text x="44" y="172" class="c">cuadros de texto o formas</text>
<text x="44" y="198" class="c">← contraportada</text>
<rect class="zona" x="324" y="46" width="262" height="80"/><text x="334" y="66" class="s">Logotipo (formas + WordArt agrupados)</text><text x="334" y="84" class="c">y nombre del restaurante</text>
<rect class="zona" x="324" y="136" width="262" height="72"/><text x="334" y="156" class="s">Imagen grande de fondo o de producto</text><text x="334" y="174" class="c">con licencia libre</text>
<text x="500" y="198" class="c">portada →</text>
<text x="20" y="252" class="t">Hoja 2 · Interior (A4 horizontal, tres columnas)</text>
<rect class="hoja" x="20" y="262" width="580" height="190"/>
<line class="pliegue" x1="213" y1="262" x2="213" y2="452"/><line class="pliegue" x1="407" y1="262" x2="407" y2="452"/>
<g class="s">
<rect class="zona" x="32" y="274" width="170" height="22"/><text x="40" y="289">PARA PICAR (título de sección)</text>
<text x="40" y="314">Plato ................ 9,50 €</text><text x="40" y="328" class="c">descripción en cursiva · 1, 3</text>
<text x="40" y="350">Plato .............. 14,00 €</text>
<rect class="zona" x="32" y="368" width="170" height="22"/><text x="40" y="383">ENSALADAS</text>
<rect class="zona" x="225" y="274" width="170" height="22"/><text x="233" y="289">CARNES</text>
<rect class="zona" x="225" y="368" width="170" height="22"/><text x="233" y="383">PESCADOS</text>
<rect class="zona" x="419" y="274" width="170" height="22"/><text x="427" y="289">POSTRES</text>
<rect class="zona" x="419" y="368" width="170" height="22"/><text x="427" y="383">BEBIDAS</text>
</g>
<text x="20" y="466" class="c">Es una distribución orientativa, no un modelo que copiar: el diseño, los colores y las imágenes los decides tú.</text>
</svg>
'''
    (IMG / 'carta-boceto.svg').write_text(svg, encoding='utf-8')


if __name__ == '__main__':
    t1_pedidos()
    t1_autorizacion()
    t2_tabla()
    t3_portada()
    t4_grafico()
    t5_invitados()
    carta_contenido()
    carta_boceto()
    print('Materiales del taller generados en', DESCARGAS, 'y', IMG)
